#!/usr/bin/env python3
# =============================================
# extractor_cli.py - 案例提取 CLI
# 用法:
#   python bin/extract.py register --name "案件名" --plaintiff "原告" --defendant "被告" ...
#   python bin/extract.py list
#   python bin/extract.py extract --case-id <id> --type defense_argument_chain --title "标题"
#   python bin/extract.py promote --artifact-id <id> --template-name "模板名" --category "答辩状"
#   python bin/extract.py complete --case-id <id>
#   python bin/extract.py search --query "关键词"
#   python bin/extract.py check --case-number "案号"
#   python bin/extract.py report
#   python bin/extract.py scan /path/to/case --dry-run
#   python bin/extract.py scan /path/to/case --ai --import
# =============================================
import argparse
import json
import os
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from src import database as db
from src import case_service as svc
from src import template_store
from src import scanner
from src import parser
from src import extractor
from src import importer


def _resolve_artifact_type(value: str) -> str:
    """解析 artifact_type：支持语义名或 D1-D5 别名。"""
    v = (value or "").strip()
    if v.upper() in db.DIMENSION_ARTIFACT_TYPES:
        return db.DIMENSION_ARTIFACT_TYPES[v.upper()]
    if v in db.ARTIFACT_TYPES:
        return v
    print(f"❌ 无效类型: {value}，有效值: D1-D5 或 {', '.join(db.ARTIFACT_TYPES)}")
    sys.exit(1)


# =============================================
# 辅助函数
# =============================================

def _print_case(c: dict) -> None:
    """打印案件概要。"""
    status_emoji = {"pending": "⏳", "extracted": "✅", "archived": "📦"}.get(
        c["status"], "❓"
    )
    n = db.count_artifacts_by_case(c["id"])
    print(
        f"  [{status_emoji}] {c['case_name']}"
        f"\n       案号: {c.get('case_number', 'N/A')}"
        f"\n       原告: {c.get('plaintiff', '')} | 被告: {c.get('defendant', '')}"
        f"\n       案由: {c.get('cause', '')}"
        f"\n       提取物: {n} 个 | 状态: {c['status']}"
        f"\n       ID: {c['id']}"
    )


def _input_multiline(prompt: str) -> str:
    """多行输入（Ctrl+D 结束）。"""
    print(f"\n{prompt}（输入完成后按 Ctrl+D 结束）:")
    lines = []
    try:
        while True:
            line = input()
            lines.append(line)
    except EOFError:
        pass
    return "\n".join(lines).strip()


def _input_required(prompt: str) -> str:
    """必填输入（不能为空）。"""
    while True:
        val = input(prompt).strip()
        if val:
            return val
        print("  ❌ 不能为空，请重新输入。")


def _select_one(options: list[str], prompt: str) -> str:
    """数字选择（单选）。"""
    for i, opt in enumerate(options, 1):
        print(f"  {i}. {opt}")
    while True:
        try:
            idx = int(input(prompt)) - 1
            if 0 <= idx < len(options):
                return options[idx]
        except ValueError:
            pass
        print("  ❌ 无效选择，请重新输入。")


# =============================================
# register 命令
# =============================================

def cmd_register(args) -> None:
    """注册新案件。"""
    print(f"\n📋 注册案件")
    print("=" * 50)

    # 全部从 args 获取（支持 --name 'xxx' 或纯交互式运行）
    name = getattr(args, "name", None) or getattr(args, "case_name", None)
    if not name:
        try:
            name = _input_required("案件名称: ")
        except EOFError:
            print("❌ 交互式输入失败，请使用 --name 等参数重试")
            sys.exit(1)
    case_number = getattr(args, "case_number", None) or ""
    plaintiff = getattr(args, "plaintiff", None) or ""
    defendant = getattr(args, "defendant", None) or ""
    cause = getattr(args, "cause", None) or ""
    result = getattr(args, "result", None) or ""
    source_path = getattr(args, "source_path", None) or ""
    notes = getattr(args, "notes", None) or ""
    entry_method = getattr(args, "entry_method", None) or "CLI"
    scene = getattr(args, "scene", None) or ""
    plaintiff_auth_level = getattr(args, "plaintiff_auth_level", None) or ""
    defendant_auth_status = getattr(args, "defendant_auth_status", None) or ""

    fields = {
        "case_name": name,
        "case_number": case_number,
        "plaintiff": plaintiff,
        "defendant": defendant,
        "cause": cause,
        "result": result,
        "source_path": source_path,
        "notes": notes,
        "entry_method": entry_method,
        "scene": scene,
        "plaintiff_auth_level": plaintiff_auth_level,
        "defendant_auth_status": defendant_auth_status,
    }

    # 识别重复
    if case_number:
        r = svc.check_duplicate(fields)
        if r["exact"]:
            print(
                f"\n⚠️  案号已存在: {r['exact']['case_name']}"
                f"\n   状态: {r['exact']['status']} | ID: {r['exact']['case_id']}"
            )
            if not input("\n是否仍要注册？（y/N）: ").strip().lower() == "y":
                print("已取消。")
                return

    try:
        case = svc.register_case(fields)
        print(f"\n✅ 案件注册成功")
        _print_case(case)
    except Exception as e:
        print(f"❌ 注册失败: {e}", file=sys.stderr)
        sys.exit(1)


# =============================================
# link 命令
# =============================================

def cmd_link(args) -> None:
    """关联原始材料路径。"""
    case = db.get_case(args.case_id)
    if not case:
        print(f"❌ 案件不存在: {args.case_id}")
        sys.exit(1)

    path = args.path or input("原始文件路径: ").strip()
    if not path:
        print("❌ 路径不能为空")
        sys.exit(1)

    # 更新 source_path（走白名单 update_case，避免私写 SQL）
    db.update_case(args.case_id, source_path=path)
    print(f"✅ 路径已关联: {path}")


# =============================================
# extract 命令
# =============================================

def cmd_extract(args) -> None:
    """提取 Artifact。"""
    case = db.get_case(args.case_id)
    if not case:
        print(f"❌ 案件不存在: {args.case_id}")
        sys.exit(1)

    print(f"\n📦 提取 Artifact")
    print(f"案件: {case['case_name']}")
    print("=" * 50)

    # 确认类型（支持语义名或 D1-D5 别名）
    if args.type:
        artifact_type = _resolve_artifact_type(args.type)
    else:
        artifact_type = _select_one(
            [f"{t} · {db.ARTIFACT_TYPE_NAMES.get(t, '')}" for t in db.ARTIFACT_TYPES],
            "选择提取物类型（数字）: ",
        ).split(" · ")[0]

    # 标题
    title = args.title or _input_required("提取物标题: ")

    # 来源引用
    source_ref = args.source_ref or input("来源引用（如：从答辩状第3段提取，可选）: ").strip()

    # 内容
    if args.content:
        content = args.content
    else:
        content_map = {
            "case_key_facts": "## 1.1 当事人信息\n\n## 1.2 案件背景\n\n## 1.3 争议焦点\n\n## 1.4 关键事实清单\n\n## 1.5 权利基础分析\n\n## 1.6 行为定性\n",
            "legal_reasoning_chain": "## 2.1 适用法律\n\n## 2.2 法律关系定性\n\n## 2.3 权利基础分析\n\n## 2.4 被告抗辩的法律可行性\n",
            "defense_argument_chain": "## 3.1 原告诉请\n\n## 3.2 被告抗辩（按优先级排序）\n\n## 3.3 抗辩论证结构\n\n## 3.4 主战场决策\n",
            "evidence_structure": "## 4.1 原告证据清单\n\n## 4.2 被告证据清单\n\n## 4.3 证据质证要点\n\n## 4.4 举证责任分配\n",
            "liability_calculation": "## 5.1 赔偿请求\n\n## 5.2 被告抗辩理由\n\n## 5.3 计算方式分析\n\n## 5.4 裁量因素\n\n## 5.5 预估结论\n",
        }
        content = _input_multiline(
            content_map.get(artifact_type, "## 内容\n\n")
        )

    artifact = svc.create_artifact(
        case_id=args.case_id,
        artifact_type=artifact_type,
        title=title,
        content=content,
        source_ref=source_ref,
    )
    print(f"\n✅ Artifact 创建成功")
    print(f"  ID: {artifact['id']}")
    print(f"  类型: {artifact['artifact_type']}")
    print(f"  标题: {artifact['title']}")
    print(f"  文件: {artifact.get('file_path', 'N/A')}")


# =============================================
# promote 命令
# =============================================

def cmd_promote(args) -> None:
    """将 Artifact 归纳为通用 Template。"""
    artifact = db.get_artifact(args.artifact_id)
    if not artifact:
        print(f"❌ 提取物不存在: {args.artifact_id}")
        sys.exit(1)

    print(f"\n📄 从 Artifact 归纳 Template")
    print(f"Artifact: {artifact['title']}（{artifact['artifact_type']}）")

    name = args.template_name or _input_required("模板名称: ")
    category = args.category or _select_one(
        db.TEMPLATE_CATEGORIES, "选择模板分类（数字）: "
    )
    template = template_store.create_template(
        name=name,
        content=artifact.get("content", ""),
        category=category,
        source_artifacts=[artifact["id"]],
        source_case_ids=[artifact["case_id"]],
    )
    print(f"\n✅ Template 创建成功")
    print(f"  ID: {template['id']}")
    print(f"  分类: {template['category']}")
    print(f"  文件: {template.get('file_path', 'N/A')}")


# =============================================
# complete 命令
# =============================================

def cmd_complete(args) -> None:
    """标记案件为 extracted。"""
    case = db.get_case(args.case_id)
    if not case:
        print(f"❌ 案件不存在: {args.case_id}")
        sys.exit(1)
    svc.update_case_status(args.case_id, "extracted")
    n = db.count_artifacts_by_case(args.case_id)
    print(f"✅ 案件标记为 extracted，共 {n} 个提取物")


# =============================================
# list 命令
# =============================================

def cmd_list(args) -> None:
    """列出案件列表。"""
    status = args.status
    cases = svc.query_cases(
        filters={"status": status, "year": args.year, "quality": args.quality}
    )
    if not cases:
        print("暂无案件")
        return
    total = len(cases)
    extracted = sum(1 for c in cases if c["status"] == "extracted")
    print(f"\n📋 案件列表（共 {total} 个，已提取 {extracted} 个）")
    print("=" * 60)
    for c in cases:
        _print_case(c)
        print()


# =============================================
# search 命令
# =============================================

def cmd_search(args) -> None:
    """搜索案件。"""
    results = svc.query_cases(search=args.query)
    if not results:
        print("未找到匹配案件")
        return
    print(f"\n🔍 搜索结果（{len(results)} 个）")
    print("=" * 60)
    for c in results:
        _print_case(c)
        print()


# =============================================
# check 命令
# =============================================

def cmd_check(args) -> None:
    """识别新案件是否重复。"""
    r = svc.check_duplicate({
        "case_number": args.case_number,
        "plaintiff": args.plaintiff or "",
        "defendant": args.defendant or "",
        "cause": args.cause or "",
    })
    print("\n🔍 案件识别结果")
    print("=" * 50)
    if r["exact"]:
        c = r["exact"]
        print(f"✅ 完全匹配（案号精确）")
        print(f"   案件: {c['case_name']}")
        print(f"   案号: {c['case_number']}")
        print(f"   状态: {c['status']} | ID: {c['case_id']}")
    elif r["partial"]:
        print(f"⚠️  当事人匹配（{len(r['partial'])} 个）")
        for c in r["partial"]:
            print(f"   - {c['case_name']} | 案由: {c['cause']} | 状态: {c['status']}")
            print(f"     ID: {c['case_id']}")
    elif r["cause_fuzzy"]:
        print(f"🔎 案由相似（{len(r['cause_fuzzy'])} 个）")
        for c in r["cause_fuzzy"]:
            print(f"   - {c['case_name']} | 案由: {c['cause']} | 置信度: {c['confidence']}")
    else:
        print("🆕 新案件（未找到匹配）")



# =============================================
# health 命令（案件健康度检查）
# =============================================

def cmd_health(args) -> None:
    """检查案件信息完整度（健康度口径同源：case_service.get_case_detail）"""
    detail = svc.get_case_detail(args.case_id)
    if not detail:
        print(f"❌ 案件不存在: {args.case_id}")
        sys.exit(1)

    case = detail["case"]
    health = detail["health"]

    field_values = {
        "案号": case.get("case_number") or "缺失",
        "原告": (case.get("plaintiff") or "缺失")[:30],
        "被告": (case.get("defendant") or "缺失")[:30],
        "案由": case.get("cause") or "缺失",
        "法院": case.get("court") or "缺失",
    }

    print(f"\n🏥 案件健康度检查")
    print("=" * 50)
    print(f"案件: {case['case_name']}")
    print(f"ID: {case['id']}")
    print()

    for mark, label in health["checks"]:
        if label.startswith("提取物"):
            print(f"  {mark} {label}")
        else:
            print(f"  {mark} {label}: {field_values.get(label, '')}")

    pct = health["health"]
    print()
    print(f"健康度: {pct}%")
    print(f"状态: {'🔴 需要补充' if pct < 50 else '🟡 基本完整' if pct < 80 else '🟢 良好'}")


# =============================================
# report 命令
# =============================================

def cmd_report(args) -> None:
    """生成案件处理报告。"""
    s = svc.summary()
    print("\n📊 案例库统计报告")
    print("=" * 50)
    print(f"  案件总数:   {s['total_cases']}")
    print(f"  已提取:     {s['extracted_cases']} 个")
    print(f"  待提取:     {s['pending_cases']} 个")
    print(f"  提取物总数: {s['total_artifacts']} 个")
    print(f"  模板总数:   {s['total_templates']} 个")

    # 按类型统计提取物
    conn = db._get_conn()
    rows = conn.execute(
        """
        SELECT artifact_type, COUNT(*) as cnt
        FROM artifacts
        GROUP BY artifact_type
        ORDER BY cnt DESC
        """
    ).fetchall()
    conn.close()
    if rows:
        print("\n  提取物类型分布:")
        for r in rows:
            print(f"    {r[0]}: {r[1]} 个")

    # 按案由统计
    conn = db._get_conn()
    cause_rows = conn.execute(
        """
        SELECT cause, COUNT(*) as cnt
        FROM cases
        WHERE cause IS NOT NULL AND cause != ''
        GROUP BY cause
        ORDER BY cnt DESC
        LIMIT 10
        """
    ).fetchall()
    if cause_rows:
        print("\n  按案由分布（Top 10）:")
        for r in cause_rows:
            bar = '█' * min(r[1], 20)
            print(f"    {r[0][:15]:<15} {bar} {r[1]}")
    conn.close()
    
    # 待提取案件
    pending = db.list_cases(status="pending")
    if pending:
        print(f"\n  待提取案件（{len(pending)} 个）:")
        for c in pending[:5]:
            print(f"    - {c['case_name'][:30]} | {c.get('case_number', 'N/A')}")



# =============================================
# export 命令（导出文书）
# =============================================

def cmd_export(args) -> None:
    """导出文书为指定格式"""
    from docx import Document
    
    # 获取 artifact
    artifact = db.get_artifact(args.artifact_id)
    if not artifact:
        print(f"❌ 提取物不存在: {args.artifact_id}")
        sys.exit(1)
    
    # 确定输出格式
    output_path = Path(args.output)
    if output_path.suffix.lower() == '.docx':
        # 导出为 DOCX
        doc = Document()
        
        # 标题
        doc.add_heading(artifact['title'], 0)
        
        # 内容（解析 Markdown）
        content = artifact.get('content', '')
        lines = content.split('\n')
        for line in lines:
            line = line.strip()
            if not line:
                doc.add_paragraph()
            elif line.startswith('# '):
                doc.add_heading(line[2:], 1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], 2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], 3)
            elif line.startswith('- ') or line.startswith('* '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('| '):
                # 表格行
                continue  # 简化处理
            else:
                doc.add_paragraph(line)
        
        doc.save(str(output_path))
        print(f"✅ 已导出为 DOCX: {output_path}")
    
    elif output_path.suffix.lower() == '.md':
        # 导出为 Markdown
        case = db.get_case(artifact['case_id'])
        frontmatter = f"""---
title: {artifact['title']}
case: {case['case_name'] if case else 'N/A'}
type: {artifact['artifact_type']}
created: {artifact['created_at']}
---

"""
        full_content = frontmatter + artifact.get('content', '')
        output_path.write_text(full_content, encoding='utf-8')
        print(f"✅ 已导出为 Markdown: {output_path}")
    
    else:
        print(f"❌ 不支持的格式: {output_path.suffix}")
        print("支持的格式: .docx, .md")


# =============================================
# templates 命令
# =============================================

def cmd_templates(args) -> None:
    """列出或搜索模板。"""
    if args.templates_action == "list":
        templates = svc.list_templates(
            args.category if args.category else None
        )
        if not templates:
            print("暂无语板")
            return
        # 按分类分组
        from collections import defaultdict
        by_cat = defaultdict(list)
        for t in templates:
            by_cat[t["category"]].append(t)
        print(f"\n📄 模板列表（共 {len(templates)} 个）")
        print("=" * 60)
        for cat, ts in sorted(by_cat.items()):
            print(f"\n  [{cat}] ({len(ts)} 个)")
            for t in ts:
                src = json.loads(t.get("source_case_ids", "[]"))
                case_hint = f"来源案件: {len(src)}个" if src else ""
                print(
                    f"    - {t['name']}  "
                    f"(v{t['version']} | 使用:{t['usage_count']}次) {case_hint}"
                )
                print(f"      ID: {t['id']}")
    elif args.templates_action == "search":
        results = svc.search_templates(args.query)
        if not results:
            print("未找到匹配模板")
            return
        print(f"\n🔍 模板搜索结果（{len(results)} 个）")
        print("=" * 60)
        for t in results:
            print(
                f"  [{t['category']}] {t['name']} "
                f"(v{t['version']} | 使用:{t['usage_count']}次)"
            )
            # 显示匹配片段
            if args.show_content:
                content = t.get("content", "")[:200]
                print(f"    内容预览: {content}...")
            print(f"    ID: {t['id']}")
    elif args.templates_action == "show":
        tmpl = db.get_template(args.template_id)
        if not tmpl:
            print(f"❌ 模板不存在: {args.template_id}")
            sys.exit(1)
        print(f"\n📄 {tmpl['name']} [{tmpl['category']}]")
        print("=" * 60)
        print(f"  ID:     {tmpl['id']}")
        print(f"  版本:   {tmpl['version']}")
        print(f"  分类:   {tmpl['category']}")
        print(f"  引用数: {tmpl['usage_count']}")
        src_cases = json.loads(tmpl.get("source_case_ids", "[]"))
        print(f"  来源案件: {', '.join(src_cases) if src_cases else '无'}")
        src_artifacts = json.loads(tmpl.get("source_artifacts", "[]"))
        print(f"  来源提取物: {', '.join(src_artifacts) if src_artifacts else '无'}")
        print()
        print(tmpl.get("content", "(无内容)"))


# =============================================
# artifacts 命令
# =============================================

def cmd_artifacts(args) -> None:
    """列出某案件下所有提取物。"""
    artifacts = svc.list_artifacts(args.case_id)
    if not artifacts:
        print("暂无提取物")
        return
    case = db.get_case(args.case_id)
    print(f"\n📦 提取物列表（{len(artifacts)} 个）| 案件: {case['case_name'] if case else args.case_id}")
    print("=" * 60)
    for a in artifacts:
        print(
            f"  [{a['artifact_type']}] {a['title']}"
            f"\n  ID: {a['id']} | 版本: {a['version']} | 来源: {a.get('source_ref', 'N/A')}"
        )


def cmd_match(args) -> None:
    """列出与某案件相似的案件（规则打分）。"""
    case = db.get_case(args.case_id)
    if not case:
        print(f"❌ 案件不存在: {args.case_id}")
        sys.exit(1)
    matches = svc.match_cases(args.case_id, limit=args.limit)
    print(f"\n🔗 相似案件（基准: {case['case_name']} | 案由: {case.get('cause') or '未知'}）")
    print("=" * 60)
    if not matches:
        print("未找到相似案件")
        return
    for i, m in enumerate(matches, 1):
        print(f"  {i}. [{m['score']}分] {m['case_name']}（{m.get('case_number') or '无案号'}）")
        print(f"     案由: {m.get('cause') or '未知'} | 年份: {m.get('year') or '未知'} | ID: {m['id']}")
        if m.get("reasons"):
            print(f"     依据: {'；'.join(m['reasons'])}")


# =============================================
# import 命令（从JSON增量导入）
# =============================================

def cmd_import(args) -> None:
    """从JSON文件增量导入案件。"""
    json_path = Path(args.json_path).resolve()
    if not json_path.exists():
        print(f"❌ 文件不存在: {json_path}")
        sys.exit(1)
    
    print(f"\n📥 增量导入: {json_path}")
    print("=" * 50)
    
    try:
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
    except json.JSONDecodeError as e:
        print(f"❌ JSON 解析失败: {e}")
        sys.exit(1)
    
    # 支持单案件和多案件
    cases = data if isinstance(data, list) else [data]
    
    success = 0
    skipped = 0
    errors = []
    
    total = len(cases)
    for i, case_data in enumerate(cases, 1):
        # 计算进度
        progress = int((i / total) * 100)
        bar_len = 30
        filled = int((i / total) * bar_len)
        bar = '█' * filled + '░' * (bar_len - filled)
        
        # 打印进度条（同一行）
        case_name = case_data.get('case_name', '未命名')[:20]
        print(f"\r[{bar}] {progress}% | {i}/{total} | {case_name}", end='', flush=True)
        
        # 构建 ExtractedCase（共用解析器，保证字段与 AI 提取一致）
        extracted = extractor.extracted_from_dict(case_data)
        
        # 增量导入（跳过已存在）
        if extracted.case_number:
            existing = db.get_case_by_number(extracted.case_number)
            if existing:
                print(f"   ⏭️  跳过（案号已存在）: {existing['case_name']}")
                skipped += 1
                continue
        
        try:
            imp = importer.CaseImporter()
            result = imp.import_full(extracted, generation_method="import-json")
            print(f"   ✅ 导入成功 (ID: {result['case_id']})")
            success += 1
        except Exception as e:
            print(f"   ❌ 导入失败: {e}")
            errors.append(f"{case_data.get('case_name')}: {e}")
    
    print(f"\n{'=' * 50}")
    print(f"📊 导入完成: 成功 {success}, 跳过 {skipped}, 失败 {len(errors)}")
    if errors:
        print("\n失败列表:")
        for e in errors:
            print(f"  - {e}")


# =============================================
# scan 命令
# =============================================

def cmd_scan(args) -> None:
    """扫描目录并可选 AI 提取。"""
    path = Path(args.path).resolve()
    
    if not path.exists():
        print(f"❌ 目录不存在: {path}")
        sys.exit(1)
    
    if not path.is_dir():
        print(f"❌ 不是目录: {path}")
        sys.exit(1)
    
    print(f"\n📁 扫描目录: {path}")
    print("⏳ 扫描中...")
    
    # Step 1: 扫描目录
    result = scanner.scan_directory(path, recursive=args.recursive, max_depth=args.max_depth)
    
    # 输出扫描结果
    print()
    print(scanner.format_scan_result(result))
    
    if args.dry_run:
        print("\n✅ 扫描完成（dry-run 模式，未解析文件内容）")
        return
    
    # Step 2: 解析文件内容
    print("\n⏳ 解析文件内容...")
    parsed_contents = parser.parse_files(result.files)
    
    if not parsed_contents:
        print("⚠️  未找到可解析的文件")
        return
    
    # 更新文件预览
    for pc in parsed_contents:
        for f in result.files:
            if f.path == pc.file_info.path:
                f.content_preview = pc.preview
                break
    
    print(f"✅ 解析完成，共 {len(parsed_contents)} 个文件")
    
    # Step 2.5: 检查案件材料完整性
    file_types = [f.category.value if hasattr(f.category, 'value') else str(f.category) for f in result.files]
    has_plaint = any('起诉状' in ft or '诉状' in ft or 'plaint' in ft.lower() for ft in file_types)
    has_defense = any('答辩' in ft or 'defense' in ft.lower() for ft in file_types)
    has_evidence = any('证据' in ft or 'evidence' in ft.lower() for ft in file_types)
    
    if not has_plaint and not has_defense and not has_evidence:
        print("\n💡 提示: 未发现标准案件文书（起诉状/答辩状/证据），建议:")
        print("   - 检查目录是否包含正确的案件材料")
        print("   - 如需 AI 提取，请使用 --ai 参数")
    
    if not args.ai:
        print("\n✅ 扫描完成（未启用 AI 提取）")
        print("   如需 AI 提取，请添加 --ai 参数")
        return
    
    # Step 3: AI 提取
    print(f"\n🤖 AI 提取中... (使用 {args.engine} 引擎)")
    print("=" * 50)
    
    try:
        extracted = extractor.extract_case(
            contents=parsed_contents,
            case_name=result.case_name,
            engine_type=args.engine,
        )
        
        # Step 3.5: 检查重复
        if extracted.case_number:
            existing = db.get_case_by_number(extracted.case_number)
            if existing:
                print(f"\n⚠️  案号已存在: {extracted.case_number}")
                print(f"   现有案件: {existing['case_name']}")
                print(f"   状态: {existing['status']} | ID: {existing['id']}")
                
                if args.yes:
                    print("\n⚠️  案号重复，跳过入库（--yes 参数）")
                    return
                
                print("\n请选择操作:")
                print("   [s] 跳过此案件")
                print("   [u] 更新现有案件")
                print("   [q] 退出")
                
                while True:
                    choice = input("\n> ").strip().lower()
                    if choice == 's':
                        print("已跳过")
                        return
                    elif choice == 'u':
                        print("将更新现有案件")
                        break
                    elif choice == 'q':
                        print("已退出")
                        return
        
        # 输出提取结果摘要
        print("\n✅ AI 提取完成")
        print("=" * 50)
        print(importer.preview_import(extracted))
        
        # 预览 Markdown
        print("\n" + "=" * 50)
        print("📄 提取结果预览 (Markdown):")
        print("=" * 50)
        print(extracted.to_markdown()[:1000] + "..." if len(extracted.to_markdown()) > 1000 else extracted.to_markdown())
        
        if not args.import_result:
            print("\n⚠️  未入库（使用 --import 参数确认入库）")
            return
        
        # Step 4: 入库
        print("\n" + "=" * 50)
        if not args.yes:
            print("⚠️  确认入库?")
            print("   [y] 确认入库")
            print("   [n] 仅保存草稿，稍后确认")
            print("   [q] 退出")
            
            while True:
                choice = input("\n> ").strip().lower()
                if choice == 'y':
                    break
                elif choice == 'n':
                    print("已保存草稿（JSON 输出到当前目录）")
                    with open(f"extracted_{result.case_name}.json", 'w', encoding='utf-8') as f:
                        json.dump(extracted.to_dict(), f, ensure_ascii=False, indent=2)
                    return
                elif choice == 'q':
                    print("已取消")
                    return
        else:
            print("⚠️  自动确认入库（--yes 参数）")
        
        # 执行入库
        gen_method = svc.ENGINE_GENERATION_METHOD.get(args.engine, "llm")
        model_version = os.environ.get("MINIMAX_MODEL", "") if args.engine == "minimax" else ""
        imp = importer.CaseImporter()
        result_import = imp.import_full(
            extracted, generation_method=gen_method, model_version=model_version
        )
        
        print("\n✅ 入库成功!")
        print(f"   案件 ID: {result_import['case_id']}")
        print(f"   提取物: {len(result_import['artifact_ids'])} 个")
        
    except extractor.ExtractionError as e:
        print(f"\n❌ AI 提取失败: {e}")
        sys.exit(1)
    except importer.ImportError as e:
        print(f"\n❌ 入库失败: {e}")
        sys.exit(1)


# =============================================
# 主入口
# =============================================

def main() -> None:
    parser = argparse.ArgumentParser(
        description="case-intelligence 提取 CLI",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    sub = parser.add_subparsers(dest="cmd", required=True)

    # register
    p = sub.add_parser("register", help="注册新案件")
    p.add_argument("--name", help="案件名称")
    p.add_argument("--case-number", dest="case_number", help="案号")
    p.add_argument("--plaintiff", help="原告")
    p.add_argument("--defendant", help="被告")
    p.add_argument("--cause", help="案由")
    p.add_argument("--result", help="案件结果")
    p.add_argument("--source-path", dest="source_path", help="原始文件路径")
    p.add_argument("--notes", help="备注")
    p.add_argument("--entry-method", dest="entry_method", choices=db.ENTRY_METHODS, default="CLI", help="录入方式（CLI/UI/AI/Bot）")
    p.add_argument("--scene", help="涉案场景（实体KTV/车载KTV/线上平台等）")
    p.add_argument("--plaintiff-auth-level", dest="plaintiff_auth_level", help="原告授权层级")
    p.add_argument("--defendant-auth-status", dest="defendant_auth_status", help="被告授权状态")
    p.set_defaults(func=cmd_register)

    # link
    p = sub.add_parser("link", help="关联原始材料路径")
    p.add_argument("--case-id", required=True, dest="case_id")
    p.add_argument("--path", help="原始文件路径")
    p.set_defaults(func=cmd_link)

    # extract
    p = sub.add_parser("extract", help="提取 Artifact")
    p.add_argument("--case-id", required=True, dest="case_id")
    p.add_argument("--type", help=f"提取物类型（{', '.join(db.ARTIFACT_TYPES)}）")
    p.add_argument("--title", help="提取物标题")
    p.add_argument("--source-ref", dest="source_ref", help="来源引用")
    p.add_argument("--content", help="提取物内容（Markdown）")
    p.set_defaults(func=cmd_extract)

    # promote
    p = sub.add_parser("promote", help="将 Artifact 归纳为 Template")
    p.add_argument("--artifact-id", required=True, dest="artifact_id")
    p.add_argument("--template-name", dest="template_name")
    p.add_argument("--category", choices=db.TEMPLATE_CATEGORIES)
    p.set_defaults(func=cmd_promote)

    # complete
    p = sub.add_parser("complete", help="标记案件为 extracted")
    p.add_argument("--case-id", required=True, dest="case_id")
    p.set_defaults(func=cmd_complete)

    # list
    p = sub.add_parser("list", help="列出所有案件")
    p.add_argument("--status", choices=["pending", "extracted", "archived"])
    p.add_argument("--year", type=int, help="按年度筛选（如 2024）")
    p.add_argument("--quality", choices=["A", "B", "C"], help="按质量等级筛选")
    p.set_defaults(func=cmd_list)

    # search
    p = sub.add_parser("search", help="搜索案件")
    p.add_argument("--query", required=True, help="搜索关键词")
    p.set_defaults(func=cmd_search)

    # check
    p = sub.add_parser("check", help="识别新案件是否重复")
    p.add_argument("--case-number", dest="case_number", help="案号")
    p.add_argument("--plaintiff", help="原告")
    p.add_argument("--defendant", help="被告")
    p.add_argument("--cause", help="案由")
    p.set_defaults(func=cmd_check)

    # health
    p = sub.add_parser("health", help="检查案件健康度")
    p.add_argument("case_id", help="案件ID")
    p.set_defaults(func=cmd_health)

    # report
    p = sub.add_parser("report", help="生成统计报告")
    p.set_defaults(func=cmd_report)

    # artifacts
    p = sub.add_parser("artifacts", help="列出案件下所有提取物")
    p.add_argument("--case-id", required=True, dest="case_id")
    p.set_defaults(func=cmd_artifacts)

    # match
    p = sub.add_parser("match", help="查找相似案件（规则打分）")
    p.add_argument("--case-id", required=True, dest="case_id")
    p.add_argument("--limit", type=int, default=5, help="返回条数（默认5）")
    p.set_defaults(func=cmd_match)

    # export
    p = sub.add_parser("export", help="导出提取物为 DOCX/Markdown")
    p.add_argument("artifact_id", help="提取物ID")
    p.add_argument("output", help="输出文件路径（.docx 或 .md）")
    p.set_defaults(func=cmd_export)

    # templates
    p = sub.add_parser("templates", help="模板管理（list/search/show）")
    sub_t = p.add_subparsers(dest="templates_action", required=True)
    # templates list
    p_list = sub_t.add_parser("list", help="列出所有模板")
    p_list.add_argument("--category", choices=db.TEMPLATE_CATEGORIES, help="按分类筛选")
    p_list.set_defaults(func=cmd_templates)
    # templates search
    p_search = sub_t.add_parser("search", help="搜索模板")
    p_search.add_argument("--query", required=True, help="搜索关键词")
    p_search.add_argument("--show-content", action="store_true", dest="show_content", help="显示内容预览")
    p_search.set_defaults(func=cmd_templates)
    # templates show
    p_show = sub_t.add_parser("show", help="查看模板详情")
    p_show.add_argument("--template-id", required=True, dest="template_id")
    p_show.set_defaults(func=cmd_templates)

    # import
    p = sub.add_parser("import", help="从JSON文件增量导入案件")
    p.add_argument("json_path", help="JSON文件路径")
    p.set_defaults(func=cmd_import)

    # scan
    p = sub.add_parser("scan", help="扫描目录并 AI 辅助录入案件")
    p.add_argument("path", help="案件目录路径")
    p.add_argument("--dry-run", action="store_true", help="仅扫描，不解析文件")
    p.add_argument("--recursive", action="store_true", default=True, help="递归扫描子目录（默认开启）")
    p.add_argument("--max-depth", type=int, default=3, help="最大递归深度（默认3）")
    p.add_argument("--ai", action="store_true", help="启用 AI 提取")
    p.add_argument("--engine", default="minimax", choices=list(extractor.SUPPORTED_ENGINES), help="AI 引擎（默认 minimax）")
    p.add_argument("--import", dest="import_result", action="store_true", help="确认入库")
    p.add_argument("--yes", "-y", action="store_true", help="跳过确认直接入库")
    p.set_defaults(func=cmd_scan)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
