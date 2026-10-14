# 案件详情页
import sys
from pathlib import Path

import streamlit as st

sys.path.insert(0, str(Path(__file__).parent.parent))

from ui.auth import require_auth
from ui.nav import render_sidebar
from ui.utils import get_case_health
from src import database as db
from src import artifact_store
from src import template_store

require_auth()

st.set_page_config(
    page_title="案例分析 - 案件详情",
    page_icon="📄",
    layout="wide"
)

render_sidebar()

if 'detail_case_id' not in st.session_state:
    st.error("未指定案件，请从案件列表进入")
    st.stop()

case_id = st.session_state.detail_case_id
case = db.get_case(case_id)

if not case:
    st.error("案件不存在")
    st.stop()

st.title(f"📄 {case.get('case_name', '案件详情')}")

STATUS_LABELS = {"pending": "待提取", "extracted": "已提取", "archived": "已归档"}

# ========== 基本信息 + 健康度 ==========
col1, col2, col3 = st.columns(3)
with col1:
    st.subheader("基本信息")
    st.write(f"**案号**: {case.get('case_number') or '无'}")
    st.write(f"**案由**: {case.get('cause') or '无'}")
    st.write(f"**法院**: {case.get('court') or '无'}")
    st.write(f"**年份**: {case.get('year') or '无'}")
    st.write(f"**状态**: {STATUS_LABELS.get(case.get('status'), case.get('status'))}")
    st.write(f"**质量**: {case.get('quality') or '无'}")

with col2:
    st.subheader("当事人")
    st.write(f"**原告**: {case.get('plaintiff') or '无'}")
    st.write(f"**被告**: {case.get('defendant') or '无'}")
    st.write(f"**场景**: {case.get('scene') or '无'}")
    if case.get('result'):
        st.write(f"**结果**: {case.get('result')}")
    if case.get('source_path'):
        st.write(f"**材料路径**: {case.get('source_path')}")

with col3:
    st.subheader("操作")
    h = get_case_health(case_id)
    st.metric("健康度", f"{h['health']}%", help=f"{h['passed']}/{h['total']} 项通过")
    for mark, label in h["checks"]:
        st.caption(f"{mark} {label}")
    if st.button("📝 生成文书", type="primary", use_container_width=True):
        st.session_state.generate_case_id = case_id
        st.switch_page("pages/4_文书生成.py")
    if st.button("🔗 相似案件", use_container_width=True):
        st.session_state.match_case_id = case_id
        st.switch_page("pages/7_案件匹配.py")
    if st.button("📋 返回列表", use_container_width=True):
        st.switch_page("pages/2_案件列表.py")

st.divider()

# ========== 编辑 / 关联 / 删除 ==========
tab_edit, tab_link, tab_danger = st.tabs(["✏️ 编辑信息", "🔗 关联材料", "🗑️ 删除案件"])

with tab_edit:
    with st.form("edit_case_form"):
        e1, e2 = st.columns(2)
        with e1:
            v_name = st.text_input("案件名称", value=case.get('case_name') or "")
            v_number = st.text_input("案号", value=case.get('case_number') or "")
            v_cause = st.text_input("案由", value=case.get('cause') or "")
            v_court = st.text_input("法院", value=case.get('court') or "")
            v_status = st.selectbox(
                "状态",
                options=["pending", "extracted", "archived"],
                index=["pending", "extracted", "archived"].index(case.get('status') or "pending")
                if (case.get('status') or "pending") in ("pending", "extracted", "archived") else 0,
                format_func=lambda x: STATUS_LABELS.get(x, x),
            )
        with e2:
            v_plaintiff = st.text_input("原告", value=case.get('plaintiff') or "")
            v_defendant = st.text_input("被告", value=case.get('defendant') or "")
            v_year = st.number_input("年份", min_value=0, max_value=2100,
                                     value=int(case.get('year') or 0), step=1)
            v_scene = st.text_input("涉案场景", value=case.get('scene') or "")
            v_quality = st.text_input("质量等级", value=case.get('quality') or "")
        v_result = st.text_input("案件结果", value=case.get('result') or "")
        v_notes = st.text_area("备注", value=case.get('notes') or "", height=60)
        if st.form_submit_button("💾 保存", type="primary"):
            try:
                db.update_case(
                    case_id,
                    case_name=v_name.strip(),
                    case_number=(v_number.strip() or None),
                    plaintiff=v_plaintiff.strip(),
                    defendant=v_defendant.strip(),
                    cause=v_cause.strip(),
                    court=v_court.strip(),
                    status=v_status,
                    year=(int(v_year) or None),
                    scene=v_scene.strip(),
                    quality=v_quality.strip(),
                    result=v_result.strip(),
                    notes=v_notes.strip(),
                )
                st.success("✅ 已保存")
                st.rerun()
            except ValueError as e:
                st.error(f"保存失败：{e}")

with tab_link:
    st.caption("关联案件的原始材料目录/文件路径（不移动原文件，仅登记引用）。")
    link_path = st.text_input("原始材料路径", value=case.get('source_path') or "")
    if st.button("💾 保存路径"):
        db.update_case(case_id, source_path=link_path.strip())
        st.success("✅ 已更新材料路径")
        st.rerun()

with tab_danger:
    st.warning("删除将级联移除该案件的全部提取物（DB + 文件），不可恢复。")
    confirm = st.checkbox("我确认删除此案件及其全部提取物")
    if st.button("🗑️ 永久删除", type="primary", disabled=not confirm):
        db.delete_case(case_id)
        artifact_store.purge_case_dir(case_id)
        st.session_state.pop("detail_case_id", None)
        st.success("已删除")
        st.switch_page("pages/2_案件列表.py")

st.divider()

# ========== 标记完成 ==========
if case.get('status') != 'extracted':
    if st.button("✅ 标记为已提取", use_container_width=True):
        db.update_case_status(case_id, "extracted")
        st.rerun()

# ========== 手动提取（A3） ==========
st.subheader("✍️ 手动提取 Artifact")
with st.expander("新增提取物", expanded=False):
    with st.form("manual_artifact_form", clear_on_submit=True):
        a_type = st.selectbox(
            "提取维度",
            options=db.ARTIFACT_TYPES,
            format_func=lambda x: f"{x}（{db.ARTIFACT_TYPE_NAMES.get(x, x)}）",
        )
        a_title = st.text_input("标题")
        a_source = st.text_input("来源引用")
        a_content = st.text_area("内容（Markdown）", height=160)
        if st.form_submit_button("➕ 新增提取物", type="primary"):
            if not a_title.strip():
                st.error("标题必填")
            else:
                artifact_store.create_artifact(
                    case_id=case_id,
                    artifact_type=a_type,
                    title=a_title.strip(),
                    content=a_content,
                    source_ref=a_source.strip(),
                    generation_method="manual",
                )
                st.success("✅ 已新增提取物")
                st.rerun()

# ========== 提取物展示（D1-D5）+ 归档为模板（A4） + 导出（A8） ==========
st.subheader("📦 提取物 (D1-D5)")
artifacts = artifact_store.list_artifacts_by_case(case_id)

if artifacts:
    for art in artifacts:
        title = art.get('title', art.get('artifact_type'))
        content = art.get('content', '')
        source = art.get('source_ref', '')

        with st.expander(f"📋 {title} ({art.get('artifact_type')})"):
            st.markdown(content)
            if source:
                st.caption(f"来源: {source}")
            with st.form(f"promote_{art['id']}"):
                p_name = st.text_input("模板名称", value=title or "")
                p_cat = st.selectbox("模板分类", options=db.TEMPLATE_CATEGORIES)
                if st.form_submit_button("📦 归档为模板"):
                    template_store.create_template(
                        name=p_name.strip() or title,
                        content=content,
                        category=p_cat,
                        source_artifacts=[art['id']],
                        source_case_ids=[case_id],
                    )
                    st.success(f"✅ 已归档为模板：{p_name or title}")

    st.divider()
    st.subheader("📤 导出提取物")
    export = artifact_store.export_case_artifacts(case_id)
    merged = export.get("markdown", "")
    base_name = f"{case.get('case_number') or 'case'}_提取物汇总"

    dl1, dl2 = st.columns(2)
    with dl1:
        st.download_button(
            "📄 导出 MD", merged,
            file_name=f"{base_name}.md", mime="text/markdown",
            use_container_width=True,
        )
    with dl2:
        try:
            from docx import Document
            import tempfile
            import os as _os

            doc = Document()
            doc.add_heading(f"{case.get('case_name')} 提取物汇总", 0)
            for art in artifacts:
                doc.add_heading(f"{art.get('title')}（{art.get('artifact_type')}）", level=1)
                doc.add_paragraph(art.get('content', ''))
            tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
            tmp.close()
            doc.save(tmp.name)
            with open(tmp.name, "rb") as fh:
                data = fh.read()
            _os.unlink(tmp.name)
            st.download_button(
                "📥 导出 DOCX", data,
                file_name=f"{base_name}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
        except ImportError:
            st.caption("DOCX 导出需 python-docx（部署环境已含）")
else:
    st.info("暂无提取物。可通过「扫描录入」或上方「手动提取」新增。")