# 文书生成页
import os
import sys
import tempfile
from pathlib import Path

import streamlit as st

sys.path.insert(0, str(Path(__file__).parent.parent))

from ui.auth import require_auth
from ui.nav import render_sidebar
from ui.utils import ENGINE_OPTIONS, DEFAULT_ENGINE
from src import database as db
from src import artifact_store
from src.document_generator import DocumentGenerator, GenerationRequest, CaseContext

require_auth()

st.set_page_config(
    page_title="案例分析 - 文书生成",
    page_icon="📝",
    layout="wide"
)

render_sidebar()

st.title("📝 文书生成")

DOC_TYPES = {
    "defense": "答辩状",
    "evidence_objection": "质证意见",
    "court_statement": "当庭陈述",
    "evidence_index": "证据目录",
}

SOURCE_LABELS = {
    "template": "模板",
    "template+llm": "模板+LLM",
    "llm": "LLM",
    "langchain_llm": "LLM",
    "fallback": "兜底（无模板/引擎）",
}

cases = db.list_cases()
case_options = {c['id']: f"{c.get('case_name', '未命名')[:30]} ({c.get('case_number', '无')})" for c in cases}

col1, col2, col3 = st.columns([1, 1, 1])
with col1:
    default_case = st.session_state.get("generate_case_id")
    selected_case_id = st.selectbox(
        "📋 选择案件",
        options=list(case_options.keys()),
        index=(list(case_options.keys()).index(default_case)
               if default_case in case_options else 0) if case_options else None,
        format_func=lambda x: case_options[x],
    )
with col2:
    doc_type = st.selectbox("📄 文书类型", options=list(DOC_TYPES.keys()),
                            format_func=lambda x: DOC_TYPES[x])
with col3:
    engine = st.selectbox("🤖 AI 引擎", options=ENGINE_OPTIONS,
                          index=ENGINE_OPTIONS.index(DEFAULT_ENGINE))

use_template = st.checkbox("使用参考模板（三级降级：模板 → 模板+LLM → LLM）", value=True)

if st.button("🤖 生成文书", type="primary", use_container_width=True):
    if not selected_case_id:
        st.error("请选择案件")
    else:
        case = db.get_case(selected_case_id)
        if not case:
            st.error("案件不存在")
        else:
            with st.spinner("🤖 生成中..."):
                try:
                    artifacts = artifact_store.list_artifacts_by_case(selected_case_id)
                    art_map = {a["artifact_type"]: a.get("content", "") for a in artifacts}
                    ctx = CaseContext(
                        case_id=selected_case_id,
                        case_number=case.get('case_number') or "",
                        case_name=case.get('case_name') or "",
                        plaintiff=case.get('plaintiff') or "",
                        defendant=case.get('defendant') or "",
                        cause=case.get('cause') or "",
                        court=case.get('court') or "",
                        scene=case.get('scene') or "",
                        artifacts=art_map,
                    )
                    gen = DocumentGenerator(engine=engine)
                    templates = gen.get_reference_templates(doc_type) if use_template else []
                    result = gen.generate(GenerationRequest(
                        doc_type=doc_type,
                        case_context=ctx,
                        reference_templates=templates,
                    ))
                    if not result.success:
                        st.error(f"生成失败: {result.error}")
                    else:
                        st.session_state.generated_content = result.content
                        st.session_state.generated_case = case
                        st.session_state.generated_doc_type = doc_type
                        st.session_state.generated_source = result.source
                        st.success(f"✅ 生成成功（来源：{SOURCE_LABELS.get(result.source, result.source)}）")
                except Exception as e:
                    st.error(f"生成失败: {str(e)}")

if 'generated_content' in st.session_state:
    st.divider()
    st.subheader(f"📄 {DOC_TYPES.get(st.session_state.generated_doc_type, '文书')} 预览")
    st.caption(f"生成来源：{SOURCE_LABELS.get(st.session_state.get('generated_source'), st.session_state.get('generated_source'))}")
    st.markdown(st.session_state.generated_content)

    st.divider()
    col1, col2 = st.columns([1, 1])
    gcase = st.session_state.generated_case
    gtype = DOC_TYPES.get(st.session_state.generated_doc_type, '文书')
    fname = f"{gcase.get('case_number') or '文书'}_{gtype}"

    with col1:
        try:
            from docx import Document

            doc = Document()
            doc.add_heading(gtype, 0)
            for line in st.session_state.generated_content.splitlines():
                doc.add_paragraph(line)
            tmp = tempfile.NamedTemporaryFile(suffix=".docx", delete=False)
            tmp.close()
            doc.save(tmp.name)
            with open(tmp.name, "rb") as fh:
                data = fh.read()
            os.unlink(tmp.name)  # C5 清理临时文件
            st.download_button(
                "📥 导出为DOCX", data, file_name=f"{fname}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True,
            )
        except ImportError:
            st.caption("DOCX 导出需 python-docx（部署环境已含）")

    with col2:
        st.download_button(
            "📄 导出为MD", st.session_state.generated_content,
            file_name=f"{fname}.md", mime="text/markdown",
            use_container_width=True,
        )