"""docx_to_pdf 字体自适应测试：多平台候选探测 + FONT_DIR 环境变量 + CID fallback。"""
from types import SimpleNamespace

import docx_to_pdf


class TestFontCandidates:
    """_font_candidates() 候选探测逻辑。"""

    def _fake_sys(self, platform: str):
        """替换模块级 sys 为假对象，模拟不同平台。"""
        docx_to_pdf.sys = SimpleNamespace(platform=platform)

    def test_font_dir_priority_over_platform(self, tmp_path, monkeypatch):
        """FONT_DIR 环境变量优先于平台内置字体。"""
        (tmp_path / "aaa.ttf").write_bytes(b"fake")
        (tmp_path / "zzz.otf").write_bytes(b"fake")
        monkeypatch.setenv("FONT_DIR", str(tmp_path))
        self._fake_sys("darwin")

        cands = docx_to_pdf._font_candidates()

        assert len(cands) == 2, "只应返回 FONT_DIR 下的字体文件"
        assert cands[0][0] == str(tmp_path / "aaa.ttf"), "按文件名排序，最先"
        assert cands[1][0] == str(tmp_path / "zzz.otf")

    def test_font_dir_filters_non_font_and_hidden(self, tmp_path, monkeypatch):
        """FONT_DIR 扫描应过滤非字体扩展名与隐藏文件。"""
        (tmp_path / "a.ttf").write_bytes(b"fake")
        (tmp_path / "b.txt").write_bytes(b"fake")
        (tmp_path / ".c.ttf").write_bytes(b"fake")
        monkeypatch.setenv("FONT_DIR", str(tmp_path))

        cands = docx_to_pdf._font_candidates()

        assert len(cands) == 1
        assert cands[0][0] == str(tmp_path / "a.ttf")

    def test_font_dir_missing_falls_back_to_platform(self, tmp_path, monkeypatch):
        """FONT_DIR 指向不存在的目录时不应报错，回退平台候选。"""
        monkeypatch.setenv("FONT_DIR", str(tmp_path / "no-such-dir"))
        self._fake_sys("darwin")

        cands = docx_to_pdf._font_candidates()

        assert len(cands) == 2
        assert cands[0][0] == "/System/Library/Fonts/Supplemental/Songti.ttc"

    def test_darwin_candidates(self, monkeypatch):
        monkeypatch.delenv("FONT_DIR", raising=False)
        self._fake_sys("darwin")
        cands = docx_to_pdf._font_candidates()
        paths = [p for p, _ in cands]
        assert any(p.startswith("/System/Library/Fonts") for p in paths), "macOS 应有系统字体路径"

    def test_linux_candidates(self, monkeypatch):
        monkeypatch.delenv("FONT_DIR", raising=False)
        self._fake_sys("linux")
        cands = docx_to_pdf._font_candidates()
        assert cands, "Linux 平台应有候选字体"
        assert cands[0][0] == "/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc"
        assert all(p.startswith("/usr/share/fonts") for p, _ in cands)

    def test_windows_candidates(self, monkeypatch):
        monkeypatch.delenv("FONT_DIR", raising=False)
        self._fake_sys("win32")
        cands = docx_to_pdf._font_candidates()
        assert cands, "Windows 平台应有候选字体"
        assert all(p.startswith("C:/Windows/Fonts") for p, _ in cands)


class TestRegisterCnFont:
    """_register_cn_font() 注册与 fallback。"""

    def test_cid_fallback_when_no_candidate(self, monkeypatch):
        """候选为空时必须回退 CID 字体，保证 PDF 仍可生成。"""
        monkeypatch.setattr(docx_to_pdf, "_font_candidates", lambda: [])
        font_name, source = docx_to_pdf._register_cn_font()
        assert font_name == "STSong-Light"
        assert "CID" in source, "来源描述应说明回退到 CID 字体"

    def test_register_smoke(self):
        """模块级注册应产出合法字体名与来源（不依赖具体平台）。"""
        assert docx_to_pdf.CN_FONT in ("CN-Font", "STSong-Light")
        assert docx_to_pdf.CN_FONT_SOURCE, "应记录字体来源，供状态/日志展示"


class TestDocxToPdfSmoke:
    """完整链路冒烟：真实 docx 重排版成 PDF，验证 CN_FONT 注册可用。"""

    def test_generate_pdf_with_chinese_text(self, tmp_path):
        from docx import Document

        docx_path = tmp_path / "sample.docx"
        pdf_path = tmp_path / "sample.pdf"

        doc = Document()
        doc.add_paragraph("法定代表人证明")
        doc.add_paragraph("兹证明张三系本单位职工，特此证明。")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "姓名"
        table.cell(0, 1).text = "张三"
        doc.save(str(docx_path))

        docx_to_pdf.docx_to_pdf(docx_path, pdf_path)

        assert pdf_path.exists()
        assert pdf_path.stat().st_size > 0, "PDF 应非空"
        assert pdf_path.read_bytes().startswith(b"%PDF"), "应产出合法 PDF 文件头"