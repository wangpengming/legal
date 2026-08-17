"""导出端点（docs/03 §6.2：POST /export）。

- scope：all（全量）/ search（复用检索条件 q+filters）/ case（指定 ids）；
- 格式：csv（utf-8-sig 带 BOM，Excel 直接打开不乱码）/ xlsx（openpyxl，未安装则 503）；
- with_footnote：尾部追加口径说明行（与 02-search 导出口径脚注一致）；
- 列名 = CASE_FIELD_META.label（中文表头），列序 = 请求 columns（默认全 47 列）；
- 每次导出写 audit_logs（op_type=export，含 scope/行数/列数）。
"""
from __future__ import annotations

import csv as csvlib
import io
import json
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, HTTPException, Response
from sqlalchemy.orm import Session

from app.api.metadata import CASE_FIELD_META
from app.api.search import _case_basic, _resolve_case, build_case_query
from app.api.schemas import ExportRequest, audit, ok
from app.core.config import DATA_DIR
from app.core.database import get_db

router = APIRouter(tags=["export"])

_FOOTNOTE_TEXT = (
    "口径说明：命中数 = FTS5 与筛选条件交集行数；支持率/单曲判赔为公式值（口径 v1.0 见 08-metrics）；"
    "金额单位：元；时间为 UTC 毫秒时间戳"
)


def _cell(v) -> str:
    """导出单元格序列化：JSON 列 → JSON 字符串；None → 空。"""
    if v is None:
        return ""
    if isinstance(v, (dict, list)):
        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, float):
        # 整数金额不带小数点（¥1,284,600 → 1284600），小数保留 2 位
        return str(int(v)) if v == int(v) else f"{v:.2f}"
    return str(v)


def _build_csv(columns: list[str], rows: list[dict], with_footnote: bool) -> str:
    buf = io.StringIO()
    writer = csvlib.writer(buf)
    writer.writerow([CASE_FIELD_META[c]["label"] for c in columns])
    for row in rows:
        writer.writerow([_cell(row.get(c)) for c in columns])
    if with_footnote:
        writer.writerow([f"# {_FOOTNOTE_TEXT}"])
    return buf.getvalue()


def _build_xlsx(columns: list[str], rows: list[dict], with_footnote: bool) -> bytes:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失走 503
        raise HTTPException(status_code=503, detail="xlsx 需要依赖 openpyxl：pip install openpyxl") from exc
    from openpyxl.styles import Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = "案件导出"
    ws.append([CASE_FIELD_META[c]["label"] for c in columns])
    for row in rows:
        ws.append([_cell(row.get(c)) for c in columns])
    if with_footnote:
        note = ws.append([f"# {_FOOTNOTE_TEXT}"])
        for cell in ws[1]:  # 表头加粗
            cell.font = cell.font.copy(bold=True)
        if note:
            ws.cell(row=ws.max_row, column=1).alignment = Alignment(wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@router.post("/export")
def export_cases(payload: ExportRequest, db: Session = Depends(get_db)) -> Response:
    """导出案件（csv/xlsx；返回文件流，Content-Disposition 附件下载）。"""
    columns = [c for c in (payload.columns or list(CASE_FIELD_META.keys())) if c in CASE_FIELD_META]
    if not columns:
        raise HTTPException(status_code=422, detail="columns 为空或全非法")

    if payload.scope == "search":
        rows, total = build_case_query(db, payload.q, payload.filters, None, columns)
    elif payload.scope == "case":
        rows, total = [], 0
        for cid in payload.ids:
            try:
                case = _resolve_case(db, cid)
            except HTTPException:
                continue
            base = _case_basic(case)
            rows.append({c: base.get(c) for c in columns})
            total += 1
    else:  # scope == "all"
        rows, total = build_case_query(db, None, [], None, columns)

    if payload.format == "xlsx":
        content = _build_xlsx(columns, rows, payload.with_footnote)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ext = "xlsx"
    else:
        content = _build_csv(columns, rows, payload.with_footnote)
        media_type = "text/csv; charset=utf-8-sig"  # BOM 由 \ufeff 前缀提供
        content = "\ufeff" + content
        ext = "csv"
    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S")
    filename = f"litigation_cases_{stamp}.{ext}"
    audit(db, op_type="export", table_name="cases", action="export",
          detail={"scope": payload.scope, "format": payload.format,
                  "rows": len(rows), "total": total, "columns": len(columns),
                  "with_footnote": payload.with_footnote})
    return Response(
        content=content if isinstance(content, str) else content,
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )