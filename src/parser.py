# =============================================
# parser.py - 文件解析器
# 支持 docx/pdf/txt/md 格式，提取文本内容
# =============================================
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from .scanner import FileInfo, FileType


@dataclass
class ParsedContent:
    """解析后的内容"""
    file_info: FileInfo
    text: str
    word_count: int
    page_count: Optional[int] = None  # PDF 页数

    @property
    def preview(self) -> str:
        """内容预览（前500字）"""
        return self.text[:500] if self.text else ""


class ParserError(Exception):
    """解析错误"""
    pass


# =============================================
# 解析器基类
# =============================================

class BaseParser:
    """解析器基类"""
    
    supported_types: set[FileType] = set()
    
    def parse(self, file_info: FileInfo) -> ParsedContent:
        """解析文件"""
        raise NotImplementedError


# =============================================
# TXT/MD 解析器
# =============================================

class TextParser(BaseParser):
    """纯文本解析器"""
    
    supported_types = {FileType.TXT, FileType.MD}
    
    def parse(self, file_info: FileInfo) -> ParsedContent:
        """解析纯文本文件"""
        try:
            # 尝试多种编码
            for encoding in ["utf-8", "gbk", "gb2312", "utf-16"]:
                try:
                    text = file_info.path.read_text(encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
            else:
                # 最后尝试二进制读取后解码
                text = file_info.path.read_bytes().decode("utf-8", errors="ignore")
            
            # 清理空白字符
            text = "\n".join(line.strip() for line in text.splitlines())
            text = "\n".join(line for line in text.split("\n") if line)
            
            return ParsedContent(
                file_info=file_info,
                text=text,
                word_count=len(text),
            )
        except Exception as e:
            raise ParserError(f"解析文本文件失败: {e}") from e


# =============================================
# DOCX 解析器
# =============================================

class DocxParser(BaseParser):
    """Word 文档解析器"""
    
    supported_types = {FileType.DOCX}
    
    def parse(self, file_info: FileInfo) -> ParsedContent:
        """解析 Word 文档"""
        try:
            from docx import Document
            
            doc = Document(str(file_info.path))
            paragraphs = []
            
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)
            
            # 也读取表格内容
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text = cell.text.strip()
                        if text:
                            paragraphs.append(text)
            
            text = "\n".join(paragraphs)
            
            return ParsedContent(
                file_info=file_info,
                text=text,
                word_count=len(text),
            )
        except ImportError:
            raise ParserError("python-docx 未安装: pip install python-docx")
        except Exception as e:
            raise ParserError(f"解析 Word 文档失败: {e}") from e


# =============================================
# PDF 解析器
# =============================================

class PdfParser(BaseParser):
    """PDF 解析器"""
    
    supported_types = {FileType.PDF}
    
    def parse(self, file_info: FileInfo) -> ParsedContent:
        """解析 PDF 文件"""
        text = ""
        page_count = 0
        
        # 尝试 pdfplumber（优先）
        try:
            import pdfplumber
            
            with pdfplumber.open(str(file_info.path)) as pdf:
                page_count = len(pdf.pages)
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except ImportError:
            pass
        except Exception:
            pass
        
        # 如果 pdfplumber 失败，尝试 pymupdf
        if not text.strip():
            try:
                import fitz  # PyMuPDF
                
                doc = fitz.open(str(file_info.path))
                page_count = len(doc)
                for page in doc:
                    page_text = page.get_text()
                    if page_text:
                        text += page_text + "\n"
                doc.close()
            except ImportError:
                raise ParserError(
                    "PDF 解析库未安装: pip install pdfplumber 或 pip install pymupdf"
                )
            except Exception as e:
                raise ParserError(f"解析 PDF 失败: {e}") from e
        
        # 清理空白
        text = "\n".join(line.strip() for line in text.splitlines())
        text = "\n".join(line for line in text.split("\n") if line)
        
        return ParsedContent(
            file_info=file_info,
            text=text,
            word_count=len(text),
            page_count=page_count,
        )


# =============================================
# 解析器工厂
# =============================================

class ParserFactory:
    """解析器工厂"""
    
    def __init__(self):
        self._parsers: dict[FileType, BaseParser] = {
            FileType.TXT: TextParser(),
            FileType.MD: TextParser(),  # MD 也用文本解析器
            FileType.DOCX: DocxParser(),
            FileType.PDF: PdfParser(),
        }
    
    def get_parser(self, file_type: FileType) -> Optional[BaseParser]:
        """获取对应类型的解析器"""
        return self._parsers.get(file_type)
    
    def parse(self, file_info: FileInfo) -> ParsedContent:
        """解析文件"""
        parser = self.get_parser(file_info.type)
        if not parser:
            raise ParserError(f"不支持的文件类型: {file_info.type.value}")
        return parser.parse(file_info)


# =============================================
# 便捷函数
# =============================================

# 全局解析器实例
_factory = ParserFactory()


def parse_file(file_info: FileInfo) -> ParsedContent:
    """解析单个文件"""
    return _factory.parse(file_info)


def parse_files(file_infos: list[FileInfo]) -> list[ParsedContent]:
    """批量解析文件"""
    results = []
    for fi in file_infos:
        try:
            result = parse_file(fi)
            results.append(result)
        except ParserError as e:
            # 记录错误但继续处理其他文件
            print(f"⚠️  解析失败: {fi.path.name} - {e}")
    return results


def extract_key_info(text: str) -> dict:
    """从文本中提取关键信息（用于快速预览）"""
    from .scanner import extract_case_number, extract_parties, extract_cause
    
    info = {
        "case_number": extract_case_number(text),
        "plaintiff": None,
        "defendant": None,
        "cause": extract_cause(text),
    }
    
    info["plaintiff"], info["defendant"] = extract_parties(text)
    
    return info
